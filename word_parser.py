# -*- coding: utf-8 -*-
from docx import Document
from openpyxl import Workbook
import re

class WordParser:
    def __init__(self):
        # 定义题目类型的正则表达式模式
        self.patterns = {
            'question_start': r'^\s*[\d]+[\.、]',  # 匹配题目开始（数字+点或顿号）
            'option': r'^\s*[A-D][\.、]',  # 匹配选项（A-D+点或顿号）
            'answer': r'^\s*【\s*答案\s*】\s*',  # 匹配答案标记，允许后面有空格
            'analysis': r'^\s*【\s*解析\s*】\s*',  # 匹配解析标记，允许后面有空格
            'title': r'^\s*(初级会计实务考试真题及答案\d+|202\d+年初级会计实务试题\d+\.\d+\s+[上下]午批次|一、[^。]+题|本类题共\d+小题|每小题[^。]+分。|错选、不选均不得分。)',  # 匹配标题和说明
            'section_title': r'^\s*[一二三四五六七八九十]+、[^。]+'  # 匹配章节标题
        }
        
    def parse_document(self, word_file, excel_file):
        """解析Word文档并保存到Excel"""
        doc = Document(word_file)
        wb = Workbook()
        ws = wb.active
        ws.title = "题库"
        
        # 设置表头
        headers = ['题号', '题目', 'A', 'B', 'C', 'D', '答案', '解析']
        ws.append(headers)
        
        current_question = {
            'number': '',
            'content': '',
            'options': {'A': '', 'B': '', 'C': '', 'D': ''},
            'answer': '',
            'analysis': ''
        }
        
        # 定义状态枚举
        STATE_CONTENT = 'content'
        STATE_OPTIONS = 'options'
        STATE_ANSWER = 'answer'
        STATE_ANALYSIS = 'analysis'
        
        current_state = STATE_CONTENT
        current_option = None
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
                
            # 跳过标题和说明性文字
            if (re.match(self.patterns['title'], text) or 
                re.match(self.patterns['section_title'], text)):
                continue
                
            # 检查是否是新题目
            if re.match(self.patterns['question_start'], text):
                # 保存上一道题目（如果存在）
                if current_question['content']:
                    self._save_question_to_excel(ws, current_question)
                
                # 初始化新题目
                current_question = {
                    'number': text.split('.')[0].strip(),
                    'content': text[text.find('.')+1:].strip(),
                    'options': {'A': '', 'B': '', 'C': '', 'D': ''},
                    'answer': '',
                    'analysis': ''
                }
                current_state = STATE_CONTENT
                current_option = None
                
            # 检查选项
            elif re.match(self.patterns['option'], text):
                current_option = text[0]
                current_question['options'][current_option] = text[2:].strip()
                current_state = STATE_OPTIONS
                
            # 检查答案和解析（可能在同一行）
            elif re.match(self.patterns['answer'], text):
                # 检查是否在同一行包含解析
                if '解析】' in text:
                    # 处理格式1：【答案】A解析】...
                    parts = text.split('解析】')
                    answer_text = re.sub(r'^\s*【\s*答案\s*】\s*', '', parts[0]).strip()
                    current_question['answer'] = answer_text
                    current_question['analysis'] = parts[1].strip()
                    current_state = STATE_ANALYSIS
                else:
                    # 处理格式2：【答案】C
                    answer_text = re.sub(r'^\s*【\s*答案\s*】\s*', '', text).strip()
                    current_question['answer'] = answer_text
                    current_state = STATE_ANSWER
                current_option = None
                
            # 检查单独的解析
            elif re.match(self.patterns['analysis'], text):
                # 提取解析内容（去掉【解析】标记和前后空格）
                analysis_text = re.sub(r'^\s*【解析】\s*', '', text).strip()
                current_question['analysis'] = analysis_text
                current_state = STATE_ANALYSIS
                current_option = None
                
            # 处理当前状态下的内容
            else:
                if current_state == STATE_CONTENT:
                    current_question['content'] += ' ' + text
                elif current_state == STATE_OPTIONS and current_option:
                    current_question['options'][current_option] += ' ' + text
                elif current_state == STATE_ANSWER:
                    # 检查是否包含解析标记
                    if '【解析】' in text:
                        parts = text.split('【解析】')
                        current_question['answer'] += ' ' + parts[0].strip()
                        current_question['analysis'] = parts[1].strip()
                        current_state = STATE_ANALYSIS
                    else:
                        current_question['answer'] += ' ' + text
                elif current_state == STATE_ANALYSIS:
                    current_question['analysis'] += ' ' + text
        
        # 保存最后一道题目
        if current_question['content']:
            self._save_question_to_excel(ws, current_question)
        
        # 调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 保存Excel文件
        wb.save(excel_file)
    
    def _save_question_to_excel(self, worksheet, question):
        """将题目数据保存到Excel工作表中"""
        # 清理多余的空格
        question['content'] = ' '.join(question['content'].split())
        question['answer'] = ' '.join(question['answer'].split())
        question['analysis'] = ' '.join(question['analysis'].split())
        for option in question['options']:
            question['options'][option] = ' '.join(question['options'][option].split())
        
        row = [
            question['number'],
            question['content'],
            question['options']['A'],
            question['options']['B'],
            question['options']['C'],
            question['options']['D'],
            question['answer'],
            question['analysis']
        ]
        worksheet.append(row)